"""Ephemeral XLSX adapter for the shared CSV profile contract."""

import hashlib
import io
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException, InvalidFileException

MAX_BYTES = 5 * 1024 * 1024

MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_MEMBERS = 128
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_SHEETS = 20


def validate_archive(content):
    """Inspect bounded ZIP/XML before handing any content to openpyxl."""
    with ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        names = [member.filename for member in members]
        if len(members) > MAX_MEMBERS or len(set(names)) != len(names):
            raise ValueError("XLSX vượt 128 mục ZIP hoặc có mục trùng; không xử lý một phần.")
        if sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
            raise ValueError("XLSX vượt 20 MiB nội dung giải nén; không xử lý một phần.")
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            raise ValueError("Không phải gói XLSX được hỗ trợ.")
        total = 0
        for member in members:
            name = member.filename
            lowered = name.lower()
            if (
                member.flag_bits & 1
                or name.startswith("/")
                or ".." in name.split("/")
                or not lowered.endswith((".xml", ".rels"))
                or any(
                    part in lowered
                    for part in ("vba", "macro", "externallink", "embedding", "activex")
                )
            ):
                raise ValueError("XLSX chứa macro, mã hóa hoặc nội dung không hỗ trợ.")
            with archive.open(member) as stream:
                data = stream.read(MAX_EXPANDED_BYTES - total + 1)
            total += len(data)
            if total > MAX_EXPANDED_BYTES:
                raise ValueError("XLSX vượt 20 MiB nội dung giải nén.")
            text = data.decode("utf-8-sig")
            if "\x00" in text or "<!DOCTYPE" in text or "<!ENTITY" in text:
                raise ValueError("XLSX chứa XML DTD/entity hoặc encoding không hỗ trợ.")
            root = ET.fromstring(text)
            if (
                name == "xl/workbook.xml"
                and sum(node.tag.rsplit("}", 1)[-1] == "sheet" for node in root.iter()) > MAX_SHEETS
            ):
                raise ValueError("XLSX vượt 20 sheet; không xử lý một phần.")
            row_count = 0
            for element in root.iter():
                tag = element.tag.rsplit("}", 1)[-1]
                if name.startswith("xl/worksheets/"):
                    if tag == "row":
                        row_count += 1
                        if (
                            row_count > MAX_ROWS + 1
                            or not 1 <= int(element.get("r", row_count)) <= MAX_ROWS + 1
                        ):
                            raise ValueError(
                                "XLSX vượt 10.000 dòng hồ sơ/sheet; không tự cắt dữ liệu."
                            )
                        if len(element) > MAX_COLUMNS:
                            raise ValueError("XLSX vượt 100 cột.")
                    if tag == "c" and element.get("r"):
                        column, row_number = coordinate_from_string(element.get("r"))
                        if (
                            row_number > MAX_ROWS + 1
                            or column_index_from_string(column) > MAX_COLUMNS
                        ):
                            raise ValueError(
                                "XLSX vượt 10.000 dòng hoặc 100 cột; không tự cắt dữ liệu."
                            )
                if tag == "f" or (tag == "t" and (element.text or "").lstrip().startswith("=")):
                    raise ValueError(
                        "XLSX chứa formula; chỉ nhận giá trị, không thực thi công thức."
                    )
                if tag == "Relationship" and element.get("TargetMode") == "External":
                    raise ValueError("XLSX chứa liên kết ngoài không hỗ trợ.")
                if tag == "Override" and element.get("PartName") == "/xl/workbook.xml":
                    if (
                        element.get("ContentType")
                        != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
                    ):
                        raise ValueError("Chỉ nhận XLSX thường, không hỗ trợ macro/template.")


def inspect_xlsx(content, sheet=None, filename="profiles.xlsx"):
    if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
        raise ValueError("Chỉ hỗ trợ XLSX (.xlsx), không nhận XLS/XLSM hoặc macro.")
    if not isinstance(content, bytes) or not content:
        raise ValueError("XLSX trống hoặc không phải nội dung byte.")
    if len(content) > MAX_BYTES:
        raise ValueError("XLSX vượt giới hạn 5 MiB; không xử lý một phần.")
    try:
        validate_archive(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (
        BadZipFile,
        CellCoordinatesException,
        InvalidFileException,
        IndexError,
        AttributeError,
        SyntaxError,
        KeyError,
        OSError,
        TypeError,
        ET.ParseError,
        UnicodeError,
        RuntimeError,
        NotImplementedError,
    ) as error:
        raise ValueError("Không đọc được XLSX hợp lệ; chưa tạo kết quả một phần.") from error
    try:
        if not workbook.worksheets or len(workbook.worksheets) != len(workbook.sheetnames):
            raise ValueError("XLSX cần ít nhất một worksheet; không hỗ trợ chart sheet.")
        if sheet is None:
            return {
                "sheets": workbook.sheetnames,
                "input_sha256": hashlib.sha256(content).hexdigest(),
            }
        if sheet not in workbook.sheetnames:
            raise ValueError("Chọn sheet có trong XLSX.")
        worksheet = workbook[sheet]
        worksheet.reset_dimensions()
        rows = []
        for cells in worksheet.iter_rows():
            values = []
            for cell in cells:
                value = cell.value
                if cell.data_type == "e" or (
                    value is not None and type(value) not in (str, int, float)
                ):
                    raise ValueError(
                        "XLSX chỉ hỗ trợ ô chuỗi/số/trống; không nhận ngày, boolean hoặc lỗi Excel."
                    )
                values.append(value)
            rows.append(tuple(values))
        if not rows:
            raise ValueError("Sheet XLSX trống; cần tên cột và hồ sơ.")
        columns = list(rows[0])
        if (
            not columns
            or any(not isinstance(name, str) or not name.strip() for name in columns)
            or len(set(columns)) != len(columns)
        ):
            raise ValueError("Tên cột XLSX phải là chuỗi có nội dung và không được trùng.")
        records = [list(row) + [None] * max(0, len(columns) - len(row)) for row in rows[1:]]
        return {"columns": columns, "records": records}
    except (KeyError, IndexError, TypeError, ET.ParseError, OSError) as error:
        raise ValueError(
            "Không đọc được cấu trúc sheet XLSX; chưa tạo kết quả một phần."
        ) from error
    finally:
        workbook.close()


def preview_xlsx(service, content, mapping, sheet, filename="profiles.xlsx"):
    import csv

    from .snapshots import csv_bytes

    if sheet is None:
        raise ValueError("Cần chọn sheet tường minh trước khi preview.")
    inspected = inspect_xlsx(content, sheet, filename)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(inspected["columns"])
    writer.writerows(inspected["records"])
    result = service.preview_profiles(output.getvalue().encode("utf-8"), mapping)
    for row, values in zip(result["rows"], inspected["records"]):
        row.pop("csv_line", None)
        row["sheet"] = sheet
        row["source_row"] = row["row_id"] + 1
        row["raw"] = dict(zip(inspected["columns"], values))
        row["raw_values"] = list(values)
    result.update(
        sheet=sheet, mapping=dict(mapping), input_sha256=hashlib.sha256(content).hexdigest()
    )
    fields = (
        list(result["rows"][0]) if result["rows"] else ["row_id", "source_row", "status", "raw"]
    )
    result["csv"] = csv_bytes(result["rows"], fields)
    return result
